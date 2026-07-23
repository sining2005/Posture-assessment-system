from __future__ import annotations

import csv
import hashlib
import json
import shutil
import tempfile
import zipfile
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from posture_assessment.database import DatabaseManager
from posture_assessment.models import AssessmentSession, ImportBatch, User
from posture_assessment.services import (
    DuplicatePatientError,
    UserInput,
    UserService,
    ValidationError,
)


CANONICAL_FIELDS = {
    "patient_no": "患者编号",
    "name": "姓名",
    "birth_date": "出生日期",
    "sex": "性别",
    "mobile": "手机",
    "height_cm": "身高",
    "weight_kg": "体重",
    "address": "地址",
    "email": "邮箱",
    "notes": "备注",
}

HEADER_ALIASES = {
    "patient_no": "patient_no",
    "patientno": "patient_no",
    "用户编号": "patient_no",
    "患者编号": "patient_no",
    "编号": "patient_no",
    "name": "name",
    "姓名": "name",
    "birthday": "birth_date",
    "birth_date": "birth_date",
    "出生日期": "birth_date",
    "sex": "sex",
    "gender": "sex",
    "性别": "sex",
    "mobile": "mobile",
    "phone": "mobile",
    "手机": "mobile",
    "手机号": "mobile",
    "手机号码": "mobile",
    "电话": "mobile",
    "height": "height_cm",
    "height_cm": "height_cm",
    "身高": "height_cm",
    "weight": "weight_kg",
    "weight_kg": "weight_kg",
    "体重": "weight_kg",
    "address": "address",
    "地址": "address",
    "email": "email",
    "邮箱": "email",
    "notes": "notes",
    "remark": "notes",
    "备注": "notes",
}


@dataclass(frozen=True, slots=True)
class ImportPreview:
    headers: list[str]
    rows: list[dict[str, Any]]
    detected_mapping: dict[str, str]


@dataclass(frozen=True, slots=True)
class ImportErrorRow:
    row_number: int
    patient_no: str
    name: str
    field: str
    original_value: str
    message: str


@dataclass(frozen=True, slots=True)
class ImportResult:
    batch_no: str
    success_count: int
    duplicate_count: int
    error_count: int
    error_rows: list[ImportErrorRow]
    error_log_path: Path | None
    source_sha256: str


class UserImportService:
    MAX_BYTES = 20 * 1024 * 1024
    MAX_ROWS = 5_000

    def __init__(
        self, db: DatabaseManager, user_service: UserService, import_log_dir: Path
    ):
        self.db = db
        self.user_service = user_service
        self.import_log_dir = Path(import_log_dir)
        self.import_log_dir.mkdir(parents=True, exist_ok=True)

    def create_template(self, destination: Path) -> Path:
        destination = Path(destination)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "用户导入模板"
        sheet.append(list(CANONICAL_FIELDS.values()))
        sheet.append(
            [
                "",
                "测试用户",
                "1990-01-01",
                "男",
                "13800138000",
                "175",
                "70",
                "北京市海淀区",
                "",
                "",
            ]
        )
        sheet.freeze_panes = "A2"
        workbook.save(destination)
        return destination

    def preview(self, source: Path) -> ImportPreview:
        source = Path(source)
        self._validate_file(source)
        headers, raw_rows = self._read(source)
        detected = {
            header: HEADER_ALIASES.get(self._normalise_header(header), "")
            for header in headers
        }
        return ImportPreview(headers=headers, rows=raw_rows, detected_mapping=detected)

    def import_file(
        self,
        source: Path,
        mapping: dict[str, str],
        duplicate_strategy: str = "skip",
        operator_name: str = "admin",
    ) -> ImportResult:
        if duplicate_strategy not in {"skip", "fill_blank", "overwrite"}:
            raise ValueError("未知的重复用户处理策略")
        source = Path(source)
        preview = self.preview(source)
        self._validate_mapping(mapping)
        source_hash = self._sha256(source)
        batch_no = f"IMP-{datetime.now():%Y%m%d-%H%M%S-%f}"
        success_count = 0
        duplicate_count = 0
        errors: list[ImportErrorRow] = []

        for offset, raw in enumerate(preview.rows, start=2):
            canonical = {
                field: raw.get(source_header)
                for source_header, field in mapping.items()
                if field
            }
            try:
                user_input = self._to_user_input(canonical)
                patient_no = (user_input.patient_no or "").strip()
                existing = (
                    self.user_service.find_by_patient_no(patient_no)
                    if patient_no
                    else None
                )
                if existing is not None:
                    duplicate_count += 1
                    if duplicate_strategy == "skip":
                        continue
                    merged = self._merge(existing, user_input, duplicate_strategy)
                    self.user_service.update(
                        existing.id, merged, require_complete=False
                    )
                    success_count += 1
                else:
                    self.user_service.create(user_input, require_complete=False)
                    success_count += 1
            except (ValidationError, DuplicatePatientError, ValueError, TypeError) as exc:
                field, value = self._guess_error_field(canonical, str(exc))
                errors.append(
                    ImportErrorRow(
                        row_number=offset,
                        patient_no=str(canonical.get("patient_no") or ""),
                        name=str(canonical.get("name") or ""),
                        field=field,
                        original_value=value,
                        message=str(exc),
                    )
                )

        error_log = self._write_error_log(batch_no, errors) if errors else None
        with self.db.session() as session:
            session.add(
                ImportBatch(
                    batch_no=batch_no,
                    source_name=source.name,
                    source_sha256=source_hash,
                    operator_name=operator_name,
                    success_count=success_count,
                    duplicate_count=duplicate_count,
                    error_count=len(errors),
                    log_path=str(error_log) if error_log else None,
                )
            )
        return ImportResult(
            batch_no=batch_no,
            success_count=success_count,
            duplicate_count=duplicate_count,
            error_count=len(errors),
            error_rows=errors,
            error_log_path=error_log,
            source_sha256=source_hash,
        )

    def _validate_file(self, source: Path) -> None:
        if not source.exists() or not source.is_file():
            raise ValidationError("导入文件不存在")
        if source.suffix.lower() not in {".xlsx", ".csv"}:
            raise ValidationError("仅支持 XLSX 或 CSV 文件")
        if source.stat().st_size > self.MAX_BYTES:
            raise ValidationError("文件不能超过 20 MB")

    def _read(self, source: Path) -> tuple[list[str], list[dict[str, Any]]]:
        if source.suffix.lower() == ".xlsx":
            workbook = load_workbook(source, read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, None)
            if not first:
                raise ValidationError("导入文件没有表头")
            headers = [str(value or "").strip() for value in first]
            rows = [
                dict(zip(headers, row, strict=False))
                for row in iterator
                if any(value not in (None, "") for value in row)
            ]
            workbook.close()
        else:
            text = None
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    text = source.read_text(encoding=encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                raise ValidationError("CSV 编码无法识别，请使用 UTF-8")
            reader = csv.DictReader(text.splitlines())
            headers = [str(value or "").strip() for value in (reader.fieldnames or [])]
            rows = [dict(row) for row in reader if any(row.values())]
        if not headers or any(not header for header in headers):
            raise ValidationError("表头不能为空")
        if len(rows) > self.MAX_ROWS:
            raise ValidationError("单次最多导入 5,000 条记录")
        return headers, rows

    def _validate_mapping(self, mapping: dict[str, str]) -> None:
        targets = [value for value in mapping.values() if value]
        if len(targets) != len(set(targets)):
            raise ValidationError("一个系统字段不能映射多次")
        for required in ("name", "birth_date", "sex"):
            if required not in targets:
                raise ValidationError(f"缺少必需字段映射：{CANONICAL_FIELDS[required]}")

    @staticmethod
    def _normalise_header(header: str) -> str:
        return str(header).strip().lower().replace(" ", "").replace("-", "_")

    @staticmethod
    def _parse_date(value: Any) -> date:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value or "").strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
        raise ValueError("出生日期格式应为 YYYY-MM-DD")

    @staticmethod
    def _parse_float(value: Any) -> float | None:
        if value in (None, ""):
            return None
        return float(str(value).strip())

    def _to_user_input(self, values: dict[str, Any]) -> UserInput:
        sex_map = {
            "m": "男",
            "male": "男",
            "1": "男",
            "男": "男",
            "f": "女",
            "female": "女",
            "2": "女",
            "女": "女",
            "": "未说明",
            "未知": "未说明",
            "未说明": "未说明",
        }
        sex_raw = str(values.get("sex") or "").strip().lower()
        sex = sex_map.get(sex_raw)
        if sex is None:
            raise ValueError("性别必须为男、女或未说明")
        return UserInput(
            patient_no=str(values.get("patient_no") or "").strip() or None,
            name=str(values.get("name") or "").strip(),
            birth_date=self._parse_date(values.get("birth_date")),
            sex=sex,
            mobile=str(values.get("mobile") or "").strip() or None,
            height_cm=self._parse_float(values.get("height_cm")),
            weight_kg=self._parse_float(values.get("weight_kg")),
            address=str(values.get("address") or "").strip() or None,
            email=str(values.get("email") or "").strip() or None,
            notes=str(values.get("notes") or "").strip() or None,
        )

    @staticmethod
    def _merge(existing: User, incoming: UserInput, strategy: str) -> UserInput:
        values = {
            "patient_no": existing.patient_no,
            "name": existing.name,
            "sex": existing.sex,
            "birth_date": existing.birth_date,
            "mobile": existing.mobile,
            "height_cm": existing.height_cm,
            "weight_kg": existing.weight_kg,
            "address": existing.address,
            "email": existing.email,
            "notes": existing.notes,
        }
        incoming_values = asdict(incoming)
        if strategy == "fill_blank":
            for field, value in incoming_values.items():
                if values.get(field) in (None, "") and value not in (None, ""):
                    values[field] = value
        else:
            for field, value in incoming_values.items():
                if value not in (None, ""):
                    values[field] = value
        return UserInput(**values)

    @staticmethod
    def _guess_error_field(values: dict[str, Any], message: str) -> tuple[str, str]:
        lookup = {
            "出生日期": "birth_date",
            "性别": "sex",
            "手机": "mobile",
            "身高": "height_cm",
            "体重": "weight_kg",
            "邮箱": "email",
            "姓名": "name",
            "患者编号": "patient_no",
        }
        for label, field in lookup.items():
            if label in message:
                return CANONICAL_FIELDS[field], str(values.get(field) or "")
        return "整行", ""

    def _write_error_log(
        self, batch_no: str, errors: Iterable[ImportErrorRow]
    ) -> Path:
        path = self.import_log_dir / f"{batch_no}-errors.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["行号", "用户编号", "姓名", "错误字段", "原始值", "处理"],
            )
            writer.writeheader()
            for item in errors:
                writer.writerow(
                    {
                        "行号": item.row_number,
                        "用户编号": item.patient_no,
                        "姓名": item.name,
                        "错误字段": item.field,
                        "原始值": item.original_value,
                        "处理": item.message,
                    }
                )
        return path

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()


@dataclass(frozen=True, slots=True)
class ExportOptions:
    user_ids: list[int]
    modules: list[str]
    include_xlsx: bool = True
    include_csv: bool = False
    include_json: bool = True
    include_landmarks: bool = True
    include_quality: bool = True
    anonymize: bool = False
    zip_package: bool = True


@dataclass(frozen=True, slots=True)
class ExportResult:
    output_path: Path
    user_count: int
    session_count: int
    sha256: str


class UserExportService:
    def __init__(self, db: DatabaseManager, export_dir: Path):
        self.db = db
        self.export_dir = Path(export_dir)
        self.export_dir.mkdir(parents=True, exist_ok=True)

    def export(self, options: ExportOptions, destination: Path | None = None) -> ExportResult:
        if not options.user_ids:
            raise ValidationError("请至少选择一名用户")
        if not options.modules:
            raise ValidationError("请至少选择一个检测模块")
        if not any((options.include_xlsx, options.include_csv, options.include_json)):
            raise ValidationError("请至少选择一种结构化数据格式")
        with self.db.session() as session:
            users = list(
                session.scalars(
                    select(User).where(
                        User.id.in_(options.user_ids), User.deleted_at.is_(None)
                    )
                )
            )
            assessments = list(
                session.scalars(
                    select(AssessmentSession).where(
                        AssessmentSession.user_id.in_(options.user_ids),
                        AssessmentSession.assessment_type.in_(options.modules),
                    )
                )
            )
        if not users:
            raise ValidationError("所选用户不存在或已归档")

        destination_dir = Path(destination or self.export_dir)
        destination_dir.mkdir(parents=True, exist_ok=True)
        export_name = f"user-data-{datetime.now():%Y%m%d-%H%M%S}"
        with tempfile.TemporaryDirectory(dir=destination_dir) as staging_text:
            staging = Path(staging_text)
            user_rows = [self._user_row(user, options.anonymize) for user in users]
            exported_ids = {
                user.id: user_rows[index]["患者编号"] for index, user in enumerate(users)
            }
            session_rows = [
                self._session_row(item, exported_ids) for item in assessments
            ]
            created: list[Path] = []
            if options.include_xlsx:
                created.append(self._write_xlsx(staging, user_rows, session_rows))
            if options.include_csv:
                created.extend(self._write_csv(staging, user_rows, session_rows))
            if options.include_json:
                created.append(self._write_json(staging, user_rows, session_rows, options))

            manifest = {
                "exported_at": datetime.now().isoformat(timespec="seconds"),
                "user_count": len(users),
                "session_count": len(assessments),
                "modules": options.modules,
                "anonymized": options.anonymize,
                "files": {
                    item.name: UserImportService._sha256(item) for item in created
                },
            }
            manifest_path = staging / "manifest.json"
            manifest_path.write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
            )

            if options.zip_package:
                output = destination_dir / f"{export_name}.zip"
                with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                    for item in [*created, manifest_path]:
                        archive.write(item, arcname=item.name)
            else:
                output = destination_dir / export_name
                output.mkdir(parents=True, exist_ok=False)
                for item in [*created, manifest_path]:
                    shutil.copy2(item, output / item.name)
        return ExportResult(
            output_path=output,
            user_count=len(users),
            session_count=len(assessments),
            sha256=(
                UserImportService._sha256(output)
                if output.is_file()
                else "directory"
            ),
        )

    @staticmethod
    def _user_row(user: User, anonymize: bool) -> dict[str, Any]:
        patient_no = user.patient_no
        name = user.name
        mobile = user.mobile or ""
        address = user.address or ""
        email = user.email or ""
        if anonymize:
            token = hashlib.sha256(patient_no.encode("utf-8")).hexdigest()[:12]
            patient_no = f"ANON-{token}"
            name = ""
            mobile = ""
            address = ""
            email = ""
        return {
            "患者编号": patient_no,
            "姓名": name,
            "出生日期": user.birth_date.isoformat(),
            "性别": user.sex,
            "手机": mobile,
            "身高(cm)": user.height_cm,
            "体重(kg)": user.weight_kg,
            "地址": address,
            "邮箱": email,
            "档案完整": "是" if user.profile_complete else "否",
        }

    @staticmethod
    def _session_row(
        item: AssessmentSession, exported_ids: dict[int, str]
    ) -> dict[str, Any]:
        return {
            "会话编号": item.session_no,
            "用户标识": exported_ids.get(item.user_id, ""),
            "检测模块": item.assessment_type,
            "状态": item.status,
            "检测时间": item.started_at.isoformat(timespec="seconds"),
            "报告路径": item.report_path or "",
        }

    @staticmethod
    def _write_xlsx(
        staging: Path,
        users: list[dict[str, Any]],
        sessions: list[dict[str, Any]],
    ) -> Path:
        path = staging / "检测数据.xlsx"
        workbook = Workbook()
        user_sheet = workbook.active
        user_sheet.title = "用户档案"
        if users:
            user_sheet.append(list(users[0]))
            for row in users:
                user_sheet.append(list(row.values()))
        session_sheet = workbook.create_sheet("检测会话")
        if sessions:
            session_sheet.append(list(sessions[0]))
            for row in sessions:
                session_sheet.append(list(row.values()))
        workbook.save(path)
        return path

    @staticmethod
    def _write_csv(
        staging: Path,
        users: list[dict[str, Any]],
        sessions: list[dict[str, Any]],
    ) -> list[Path]:
        paths: list[Path] = []
        for name, rows in (("用户档案.csv", users), ("检测会话.csv", sessions)):
            path = staging / name
            if rows:
                with path.open("w", encoding="utf-8-sig", newline="") as handle:
                    writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
                    writer.writeheader()
                    writer.writerows(rows)
            else:
                path.write_text("", encoding="utf-8")
            paths.append(path)
        return paths

    @staticmethod
    def _write_json(
        staging: Path,
        users: list[dict[str, Any]],
        sessions: list[dict[str, Any]],
        options: ExportOptions,
    ) -> Path:
        path = staging / "原始结构.json"
        payload = {
            "users": users,
            "assessment_sessions": sessions,
            "requested_modules": options.modules,
            "include_landmarks": options.include_landmarks,
            "include_quality": options.include_quality,
        }
        path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return path
